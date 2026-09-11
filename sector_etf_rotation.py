"""
SECTOR ETF ROTATION research — FULL / local only.

Visual study of Select Sector ETF strength vs SPY over 20D / 40D / 63D.
IBKR Stock/ETF daily bars only (no Yahoo). Reuses Market Index window,
TOTAL compounding, heatmap, SOURCE/TIME, and column-pref patterns.

Main new calculation: RELATIVE TO SPY (daily and compounded).
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from typing import Any
from zoneinfo import ZoneInfo

from db import get_conn, get_setting, init_db, set_setting
from sector_etf_config import (
    BENCHMARK_STORAGE,
    DEFAULT_GROUP_KEY,
    all_storage_tickers,
    list_groups,
    list_sector_etfs,
)
from strong_stocks import upsert_daily_bars

log = logging.getLogger("leibot.sector_etf_rotation")
ET = ZoneInfo("America/New_York")

WINDOW_CHOICES = (20, 40, 63)
DEFAULT_WINDOW = 63
MODE_CHOICES = ("raw", "rel")
DEFAULT_MODE = "rel"

SETTINGS_COLUMNS_KEY = "sector_etf_rotation_columns_v3"
SETTINGS_META_KEY = "sector_etf_rotation_refresh_meta"

# Default visible: ETF / SPY / REL each show 20D + 40D + 63D.
DEFAULT_COLUMNS: dict[str, bool] = {
    "sector": True,
    "daily": True,
    "ret_5d": False,
    "ret_20d": True,
    "ret_40d": True,
    "ret_63d": True,
    "spy_5d": False,
    "spy_20d": True,
    "spy_40d": True,
    "spy_63d": True,
    "total": True,
    "rel_5d": False,
    "rel_20d": True,
    "rel_40d": True,
    "rel_63d": True,
    "rel_rank": True,
    "rank_change": True,
    "source": True,
    "time": True,
}

PERIOD_LOOKBACKS = (5, 20, 40, 63)


def get_column_prefs() -> dict[str, bool]:
    raw = get_setting(SETTINGS_COLUMNS_KEY, None)
    out = dict(DEFAULT_COLUMNS)
    if isinstance(raw, dict):
        for k in DEFAULT_COLUMNS:
            if k in raw:
                out[k] = bool(raw[k])
    return out


def set_column_prefs(prefs: dict[str, Any]) -> dict[str, bool]:
    out = dict(DEFAULT_COLUMNS)
    if isinstance(prefs, dict):
        for k in DEFAULT_COLUMNS:
            if k in prefs:
                out[k] = bool(prefs[k])
    set_setting(SETTINGS_COLUMNS_KEY, out)
    return out


def get_refresh_meta() -> dict[str, Any]:
    raw = get_setting(SETTINGS_META_KEY, None)
    return dict(raw) if isinstance(raw, dict) else {}


def _set_refresh_meta(meta: dict[str, Any]) -> None:
    set_setting(SETTINGS_META_KEY, meta)


def _fmt_time(raw: Any) -> str:
    if not raw:
        return "—"
    s = str(raw).strip()
    try:
        if s.endswith("Z"):
            s = s[:-1] + "+00:00"
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(ET).strftime("%Y-%m-%d %H:%M ET")
    except Exception:
        return s[:19] if len(s) >= 19 else s


def _load_dated_closes(
    tickers: list[str], *, lookback_calendar_days: int = 400
) -> dict[str, list[tuple[str, float]]]:
    init_db()
    clean = sorted({(t or "").strip().upper() for t in tickers if t})
    if not clean:
        return {}
    ph = ",".join("?" * len(clean))
    with get_conn() as conn:
        rows = conn.execute(
            f"""
            SELECT ticker, date, close FROM daily_bars
            WHERE ticker IN ({ph})
              AND date >= date('now', ?)
            ORDER BY ticker ASC, date ASC
            """,
            [*clean, f"-{int(lookback_calendar_days)} days"],
        ).fetchall()
    out: dict[str, list[tuple[str, float]]] = {}
    for r in rows:
        t = (r["ticker"] or "").strip().upper()
        d = str(r["date"] or "")[:10]
        try:
            px = float(r["close"])
        except (TypeError, ValueError):
            continue
        if not d or px <= 0:
            continue
        out.setdefault(t, []).append((d, px))
    return out


def _trading_calendar(
    dated: dict[str, list[tuple[str, float]]], *, window: int
) -> list[str]:
    from collections import Counter

    # Prefer SPY calendar when present; else majority of sector ETFs.
    spy = dated.get(BENCHMARK_STORAGE) or []
    if len(spy) >= window:
        return [d for d, _ in spy[-window:]]

    counts: Counter[str] = Counter()
    n = 0
    for key, series in dated.items():
        if key == BENCHMARK_STORAGE or not series:
            continue
        n += 1
        for d, _ in series[-(window + 15) :]:
            counts[d] += 1
    if not counts:
        if spy:
            return [d for d, _ in spy[-window:]]
        return []
    threshold = max(1, n // 2)
    ordered = sorted(d for d, c in counts.items() if c >= threshold)
    if len(ordered) < window:
        ordered = sorted(counts.keys())
    if len(ordered) <= window:
        return ordered
    return ordered[-window:]


def _period_return_pct(series: list[tuple[str, float]], lookback: int) -> float | None:
    if lookback < 1 or len(series) < lookback + 1:
        return None
    base = series[-(lookback + 1)][1]
    last = series[-1][1]
    if base <= 0 or last <= 0:
        return None
    return round((last / base - 1.0) * 100.0, 2)


def _period_return_pct_offset(
    series: list[tuple[str, float]], lookback: int, *, end_offset: int
) -> float | None:
    """Return ending `end_offset` bars before the latest bar."""
    if lookback < 1 or end_offset < 0:
        return None
    need = lookback + 1 + end_offset
    if len(series) < need:
        return None
    last = series[-(1 + end_offset)][1]
    base = series[-(lookback + 1 + end_offset)][1]
    if base <= 0 or last <= 0:
        return None
    return round((last / base - 1.0) * 100.0, 2)


def _compound_total(daily_pcts: list[float | None]) -> list[float | None]:
    """TOTAL starts at 100 on day 0; then compounds Daily Change %."""
    out: list[float | None] = []
    total: float | None = None
    for i, pct in enumerate(daily_pcts):
        if i == 0:
            total = 100.0
            out.append(100.0)
            continue
        if total is None or pct is None:
            out.append(None)
            continue
        total = total * (1.0 + float(pct) / 100.0)
        out.append(round(total, 2))
    return out


def _daily_pct_map(series: list[tuple[str, float]]) -> dict[str, float]:
    out: dict[str, float] = {}
    for i, (d, px) in enumerate(series):
        if i == 0:
            continue
        prior = series[i - 1][1]
        if prior > 0 and px > 0:
            out[d] = round((px / prior - 1.0) * 100.0, 2)
    return out


def refresh_sector_etfs_from_ibkr(
    *, period: str = "1y", group_key: str | None = None
) -> dict[str, Any]:
    """
    Pull IBKR ETF daily bars into daily_bars. Never uses Yahoo.
    On failure, leaves prior stored bars untouched.
    """
    tickers = all_storage_tickers(group_key)
    result: dict[str, Any] = {
        "ok": False,
        "provider": "ibkr",
        "updated": 0,
        "failed": [],
        "errors": {},
        "as_of": "",
    }
    try:
        from ibkr_local.adapter import get_adapter

        adapter = get_adapter()
        fetched = adapter.fetch_daily_bars(tickers, period=period)
    except Exception as exc:
        log.exception("sector ETF IBKR fetch failed")
        result["errors"]["__all__"] = str(exc)
        return result

    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    per_etf: dict[str, Any] = {}
    latest_dates: list[str] = []
    updated = 0
    for sym in tickers:
        payload = (fetched or {}).get(sym) or {}
        if not payload.get("ok"):
            err = payload.get("error") or "fetch failed"
            result["failed"].append(sym)
            result["errors"][sym] = err
            per_etf[sym] = {
                "ok": False,
                "error": err,
                "provider": "ibkr",
                "updated_at": now,
            }
            continue
        closes = payload.get("closes")
        if closes is None or getattr(closes, "empty", True):
            result["failed"].append(sym)
            result["errors"][sym] = "empty closes"
            per_etf[sym] = {
                "ok": False,
                "error": "empty closes",
                "provider": "ibkr",
                "updated_at": now,
            }
            continue
        n = upsert_daily_bars(sym, closes)
        updated += int(n or 0)
        bar_day = str(payload.get("latest_bar_date") or "")[:10]
        if bar_day:
            latest_dates.append(bar_day)
        per_etf[sym] = {
            "ok": True,
            "provider": "ibkr",
            "bars": int(payload.get("bars") or 0),
            "latest_bar_date": bar_day,
            "updated_at": now,
            "mode": payload.get("mode"),
        }

    as_of = max(latest_dates) if latest_dates else ""
    meta = {
        "provider": "IBKR",
        "updated_at": now,
        "as_of": as_of,
        "group_key": group_key or DEFAULT_GROUP_KEY,
        "etfs": per_etf,
    }
    _set_refresh_meta(meta)
    result.update(
        {
            "ok": updated > 0,
            "updated": updated,
            "as_of": as_of,
            "updated_at": now,
            "meta": meta,
        }
    )
    return result


def _build_etf_row(
    edef: dict[str, Any],
    *,
    dates: list[str],
    dated: dict[str, list[tuple[str, float]]],
    spy_daily: dict[str, float],
    spy_series: list[tuple[str, float]],
    meta: dict[str, Any],
    window: int,
    is_benchmark: bool = False,
) -> dict[str, Any]:
    st = edef["storage_ticker"]
    series = dated.get(st) or []
    cmap = {d: px for d, px in series}
    etf_daily = _daily_pct_map(series)

    days_out: list[dict[str, Any]] = []
    raw_pcts: list[float | None] = []
    rel_pcts: list[float | None] = []
    for d in dates:
        close = cmap.get(d)
        raw = etf_daily.get(d)
        spy = spy_daily.get(d)
        rel: float | None = None
        if is_benchmark:
            # SPY vs itself: relative move is 0 when a daily % exists.
            if raw is not None:
                rel = 0.0
        elif raw is not None and spy is not None:
            rel = round(float(raw) - float(spy), 2)
        raw_pcts.append(raw)
        rel_pcts.append(rel)
        days_out.append(
            {
                "date": d,
                "daily_pct": raw,
                "rel_pct": rel,
                "close": round(close, 4) if close is not None else None,
            }
        )

    raw_totals = _compound_total(raw_pcts)
    rel_totals = _compound_total(rel_pcts)
    for i, day in enumerate(days_out):
        day["total"] = raw_totals[i] if i < len(raw_totals) else None
        day["rel_total"] = rel_totals[i] if i < len(rel_totals) else None

    rets = {f"ret_{n}d": _period_return_pct(series, n) for n in PERIOD_LOOKBACKS}
    spy_rets = {
        f"spy_{n}d": _period_return_pct(spy_series, n) for n in PERIOD_LOOKBACKS
    }
    rels: dict[str, float | None] = {}
    for n in PERIOD_LOOKBACKS:
        if is_benchmark:
            # Benchmark relative period return is 0 when SPY return exists.
            p = spy_rets.get(f"spy_{n}d")
            rels[f"rel_{n}d"] = 0.0 if p is not None else None
        else:
            s = _period_return_pct(series, n)
            p = spy_rets.get(f"spy_{n}d")
            rels[f"rel_{n}d"] = (
                round(float(s) - float(p), 2)
                if s is not None and p is not None
                else None
            )

    # Previous-window relative return (same length, ending `window` bars ago).
    prev_rel: float | None = None
    if not is_benchmark:
        prev_s = _period_return_pct_offset(series, window, end_offset=window)
        prev_p = _period_return_pct_offset(spy_series, window, end_offset=window)
        prev_rel = (
            round(float(prev_s) - float(prev_p), 2)
            if prev_s is not None and prev_p is not None
            else None
        )

    etf_meta = (meta.get("etfs") or {}).get(st) or {}
    provider = "—"
    if series:
        provider = "IBKR"
    elif etf_meta.get("provider"):
        provider = str(etf_meta.get("provider") or "IBKR").upper()
        if provider == "IBKR" and not etf_meta.get("ok"):
            provider = "—"

    updated_at = etf_meta.get("updated_at") or meta.get("updated_at") or ""
    last_total = days_out[-1]["total"] if days_out else None
    last_rel_total = days_out[-1]["rel_total"] if days_out else None
    window_ret = round(last_total - 100.0, 2) if last_total is not None else None
    window_rel_ret = (
        round(last_rel_total - 100.0, 2) if last_rel_total is not None else None
    )

    return {
        "key": edef["key"],
        "etf": edef["etf"],
        "sector": edef["sector"],
        "storage_ticker": st,
        "is_benchmark": bool(is_benchmark),
        "days": days_out,
        "total": last_total,
        "rel_total": last_rel_total,
        "window_return_pct": window_ret,
        "window_rel_return_pct": window_rel_ret,
        **rets,
        **spy_rets,
        **rels,
        "prev_rel_score": prev_rel,
        "data_source": provider if series else "—",
        "downloaded_at": updated_at,
        "downloaded_at_display": _fmt_time(updated_at),
        "relative_rank": None,
        "prev_relative_rank": None,
        "rank_change": None,
        "strength": "Benchmark" if is_benchmark else "—",
    }


def _assign_relative_ranks(rows: list[dict[str, Any]], *, window: int) -> None:  # noqa: ARG001
    """
    RELATIVE RANK by window REL TOTAL (higher = outperforming SPY more).
    RANK CHANGE vs previous non-overlapping window of the same length.
    SPY benchmark row is excluded from ranking.
    """
    scored = [
        (i, r.get("rel_total"))
        for i, r in enumerate(rows)
        if not r.get("is_benchmark") and r.get("rel_total") is not None
    ]
    scored.sort(key=lambda x: (-float(x[1]), rows[x[0]].get("etf") or ""))
    for rank, (idx, _) in enumerate(scored, start=1):
        rows[idx]["relative_rank"] = rank

    prev_scored = [
        (i, r.get("prev_rel_score"))
        for i, r in enumerate(rows)
        if not r.get("is_benchmark") and r.get("prev_rel_score") is not None
    ]
    prev_scored.sort(key=lambda x: (-float(x[1]), rows[x[0]].get("etf") or ""))
    for rank, (idx, _) in enumerate(prev_scored, start=1):
        rows[idx]["prev_relative_rank"] = rank

    for r in rows:
        if r.get("is_benchmark"):
            r["relative_rank"] = None
            r["prev_relative_rank"] = None
            r["rank_change"] = None
            r["strength"] = "Benchmark"
            continue
        cur = r.get("relative_rank")
        prev = r.get("prev_relative_rank")
        if cur is None or prev is None:
            r["rank_change"] = None
            continue
        # Positive change = improved rank (e.g. was #5, now #1 → +4).
        change = int(prev) - int(cur)
        r["rank_change"] = change
        if change >= 2:
            r["strength"] = "Improving"
        elif change <= -2:
            r["strength"] = "Weakening"
        else:
            r["strength"] = "Neutral"

    # If no prev ranks, fall back to band vs median REL TOTAL.
    if not prev_scored and len(scored) >= 2:
        totals = [float(t) for _, t in scored]
        mid = sorted(totals)[len(totals) // 2]
        for idx, tot in scored:
            t = float(tot)
            if t >= mid + 1.0:
                rows[idx]["strength"] = "Improving"
            elif t <= mid - 1.0:
                rows[idx]["strength"] = "Weakening"
            else:
                rows[idx]["strength"] = "Neutral"


def load_sector_etf_rotation(
    *,
    window: int = DEFAULT_WINDOW,
    mode: str = DEFAULT_MODE,
    group_key: str | None = None,
    refresh: bool = False,
) -> dict[str, Any]:
    """
    Build SECTOR ETF ROTATION payload.
    refresh=False: cache only.
    refresh=True: IBKR ETF bars only; keep prior DB on failure.
    mode=raw|rel controls heatmap / primary TOTAL display.
    """
    try:
        w = int(window)
    except (TypeError, ValueError):
        w = DEFAULT_WINDOW
    if w not in WINDOW_CHOICES:
        w = min(WINDOW_CHOICES, key=lambda x: abs(x - w))

    m = (mode or DEFAULT_MODE).strip().lower()
    if m not in MODE_CHOICES:
        m = DEFAULT_MODE

    gk = (group_key or DEFAULT_GROUP_KEY).strip().lower() or DEFAULT_GROUP_KEY
    defs = list_sector_etfs(gk)
    refresh_result: dict[str, Any] | None = None
    if refresh:
        refresh_result = refresh_sector_etfs_from_ibkr(period="1y", group_key=gk)

    meta = get_refresh_meta()
    tickers = all_storage_tickers(gk)
    # Need ~2 windows for RANK CHANGE previous period.
    dated = _load_dated_closes(tickers, lookback_calendar_days=max(280, w * 8))
    dates = _trading_calendar(dated, window=w)
    spy_series = dated.get(BENCHMARK_STORAGE) or []
    spy_daily = _daily_pct_map(spy_series)

    rows = [
        _build_etf_row(
            d,
            dates=dates,
            dated=dated,
            spy_daily=spy_daily,
            spy_series=spy_series,
            meta=meta,
            window=w,
        )
        for d in defs
    ]
    _assign_relative_ranks(rows, window=w)
    rows.sort(
        key=lambda r: (
            r["relative_rank"] if r.get("relative_rank") is not None else 999,
            r.get("etf") or "",
        )
    )
    # List SPY as benchmark row at the top (not ranked with sectors).
    spy_row = _build_etf_row(
        {
            "key": "spy",
            "etf": BENCHMARK_STORAGE,
            "sector": "Market Benchmark",
            "storage_ticker": BENCHMARK_STORAGE,
        },
        dates=dates,
        dated=dated,
        spy_daily=spy_daily,
        spy_series=spy_series,
        meta=meta,
        window=w,
        is_benchmark=True,
    )
    rows = [spy_row] + rows
    # row_count for badge = sector ETFs only (exclude SPY benchmark).
    sector_count = sum(1 for r in rows if not r.get("is_benchmark"))

    date_labels = []
    for d in dates:
        try:
            _y, mo, day = d.split("-")
            months = (
                "Jan",
                "Feb",
                "Mar",
                "Apr",
                "May",
                "Jun",
                "Jul",
                "Aug",
                "Sep",
                "Oct",
                "Nov",
                "Dec",
            )
            label = f"{months[int(mo) - 1]} {int(day)}"
        except Exception:
            label = d
        date_labels.append({"date": d, "label": label})

    as_of = dates[-1] if dates else (meta.get("as_of") or "")
    spy_ok = bool(spy_series)
    return {
        "window": w,
        "window_choices": list(WINDOW_CHOICES),
        "mode": m,
        "mode_choices": list(MODE_CHOICES),
        "group_key": gk,
        "groups": list_groups(),
        "dates": dates,
        "date_labels": date_labels,
        "rows": rows,
        "row_count": sector_count,
        "columns": get_column_prefs(),
        "as_of": as_of,
        "refreshed": bool(refresh),
        "refresh_result": refresh_result,
        "benchmark": BENCHMARK_STORAGE,
        "spy_ok": spy_ok,
        "provider": meta.get("provider") or ("IBKR" if any(dated.values()) else "—"),
        "downloaded_at_display": _fmt_time(meta.get("updated_at")),
        "notes": (
            "RAW shows ETF Daily Change %. RELATIVE TO SPY shows ETF% − SPY%. "
            "REL TOTAL compounds relative daily moves from 100. "
            "Relative Rank orders by REL TOTAL for the selected window. "
            "Rank Change compares vs the prior non-overlapping window. "
            "Relative strength ≠ literal money flow. IBKR only — no Yahoo."
        ),
    }


def export_sector_etf_rotation_xlsx(
    *, window: int = DEFAULT_WINDOW, mode: str = DEFAULT_MODE, group_key: str | None = None
) -> tuple[str, bytes]:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    payload = load_sector_etf_rotation(
        window=window, mode=mode, group_key=group_key, refresh=False
    )
    dates = list(payload.get("dates") or [])
    rows = list(payload.get("rows") or [])
    as_of = payload.get("as_of") or "na"
    mode_s = payload.get("mode") or mode
    fname = (
        f"sector_etf_rotation_{payload.get('window') or window}D_"
        f"{mode_s}_{as_of}.xlsx"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Sector ETF Rotation"
    header = [
        "ETF",
        "Sector",
        "Rel Rank",
        "Prev Rank",
        "Rank Change",
        "Strength",
        "5D %",
        "20D %",
        "40D %",
        "63D %",
        "SPY 5D %",
        "SPY 20D %",
        "SPY 40D %",
        "SPY 63D %",
        "TOTAL",
        "REL 5D %",
        "REL 20D %",
        "REL 40D %",
        "REL 63D %",
        "REL TOTAL",
        "Source",
        "Time",
    ]
    for d in dates:
        header.extend([f"{d} Daily %", f"{d} Rel %", f"{d} Close", f"{d} TOTAL", f"{d} Rel TOTAL"])
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        line: list[Any] = [
            r.get("etf") or "",
            r.get("sector") or "",
            r.get("relative_rank"),
            r.get("prev_relative_rank"),
            r.get("rank_change"),
            r.get("strength") or "",
            r.get("ret_5d"),
            r.get("ret_20d"),
            r.get("ret_40d"),
            r.get("ret_63d"),
            r.get("spy_5d"),
            r.get("spy_20d"),
            r.get("spy_40d"),
            r.get("spy_63d"),
            r.get("total"),
            r.get("rel_5d"),
            r.get("rel_20d"),
            r.get("rel_40d"),
            r.get("rel_63d"),
            r.get("rel_total"),
            r.get("data_source") or "",
            r.get("downloaded_at_display") or "",
        ]
        by_date = {d.get("date"): d for d in (r.get("days") or []) if d.get("date")}
        for d in dates:
            cell = by_date.get(d) or {}
            line.extend(
                [
                    cell.get("daily_pct"),
                    cell.get("rel_pct"),
                    cell.get("close"),
                    cell.get("total"),
                    cell.get("rel_total"),
                ]
            )
        ws.append(line)

    for i, width in enumerate(
        (8, 22, 8, 8, 10, 12, 8, 8, 8, 8, 9, 9, 9, 9, 10, 9, 9, 9, 9, 10, 8, 22),
        start=1,
    ):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    bio = io.BytesIO()
    wb.save(bio)
    return fname, bio.getvalue()
