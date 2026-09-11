"""
Group Movement (Herd Analysis) — Full / local research only.

Visual study of NASDAQ-100 group co-movement over a recent trading window.
No Leader/Follower labels, scores, or trade signals.
"""

from __future__ import annotations

import logging
from collections import Counter
from datetime import datetime, time, timedelta, timezone
from typing import Any
from zoneinfo import ZoneInfo

from db import get_conn, get_setting, init_db, list_universe, set_setting
from sector_rotation import normalize_sector_name

log = logging.getLogger("leibot.group_movement")
ET = ZoneInfo("America/New_York")

WINDOW_CHOICES = (20, 40, 63)
DEFAULT_WINDOW = 40
UNCLASSIFIED = "Unclassified"
ALL_GROUP_KEY = "ALL"

# Soft UI prefs (hiding columns never stops collection).
SETTINGS_COLUMNS_KEY = "group_movement_columns"
DEFAULT_COLUMNS: dict[str, bool] = {
    "daily": True,
    "pre": False,
    "regular": False,
    "after": False,
    "price": False,
}


def get_column_prefs() -> dict[str, bool]:
    raw = get_setting(SETTINGS_COLUMNS_KEY, None)
    out = dict(DEFAULT_COLUMNS)
    if isinstance(raw, dict):
        for k in DEFAULT_COLUMNS:
            if k in raw:
                out[k] = bool(raw[k])
    return out


def set_column_prefs(prefs: dict[str, Any]) -> dict[str, bool]:
    out = dict(DEFAULT_COLUMNS)
    if isinstance(prefs, dict):
        for k in DEFAULT_COLUMNS:
            if k in prefs:
                out[k] = bool(prefs[k])
    set_setting(SETTINGS_COLUMNS_KEY, out)
    return out


# When universe.sector is blank, map common industry labels → GICS sector.
_INDUSTRY_SECTOR_HINTS: tuple[tuple[str, str], ...] = (
    ("biotechnology", "Health Care"),
    ("pharmaceutical", "Health Care"),
    ("health care equipment", "Health Care"),
    ("health care", "Health Care"),
    ("healthcare", "Health Care"),
    ("medical", "Health Care"),
    ("semiconductor", "Information Technology"),
    ("software", "Information Technology"),
    ("technology", "Information Technology"),
    ("information technology", "Information Technology"),
    ("telecommunications", "Communication Services"),
    ("communication", "Communication Services"),
    ("media", "Communication Services"),
    ("consumer discretionary", "Consumer Discretionary"),
    ("consumer staples", "Consumer Staples"),
    ("industrials", "Industrials"),
    ("industrial", "Industrials"),
    ("financial", "Financials"),
    ("bank", "Financials"),
    ("utilities", "Utilities"),
    ("utility", "Utilities"),
    ("energy", "Energy"),
    ("materials", "Materials"),
    ("real estate", "Real Estate"),
)


def _canon_sector(raw: str | None, industry: str | None = None) -> str:
    n = normalize_sector_name(raw)
    if n:
        return n
    s = (raw or "").strip()
    if s:
        # Already a usable label (even if not in GICS map).
        return s

    ind = (industry or "").strip()
    if ind and ind != "—":
        n2 = normalize_sector_name(ind)
        if n2:
            return n2
        key = ind.lower()
        for needle, sector in _INDUSTRY_SECTOR_HINTS:
            if key == needle or needle in key:
                return sector
        # Exact industry text sometimes equals a GICS name.
        for needle, sector in _INDUSTRY_SECTOR_HINTS:
            if key == sector.lower():
                return sector
    return UNCLASSIFIED


def list_ndx100_members() -> list[dict[str, Any]]:
    """NASDAQ-100 universe rows with normalized sector + industry."""
    rows = list_universe(group="ndx100") or []
    out: list[dict[str, Any]] = []
    for r in rows:
        t = (r.get("ticker") or "").strip().upper()
        if not t:
            continue
        industry = (r.get("industry") or "").strip() or "—"
        sector = _canon_sector(r.get("sector"), industry)
        out.append(
            {
                "ticker": t,
                "name": (r.get("name") or "").strip(),
                "sector": sector,
                "industry": industry,
            }
        )
    out.sort(key=lambda x: (x["sector"], x["industry"], x["ticker"]))
    return out


def list_movement_groups(members: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """
    Selectable groups: ALL (full NASDAQ-100), each Sector, plus Sector||Industry slices.
    key format: ALL  OR  sector  OR  sector||industry
    """
    members = members if members is not None else list_ndx100_members()
    by_sector: dict[str, list[dict[str, Any]]] = {}
    by_ind: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for m in members:
        by_sector.setdefault(m["sector"], []).append(m)
        by_ind.setdefault((m["sector"], m["industry"]), []).append(m)

    groups: list[dict[str, Any]] = [
        {
            "key": ALL_GROUP_KEY,
            "label": "ALL",
            "kind": "all",
            "count": len(members),
        }
    ]
    for sector, rows in sorted(by_sector.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        groups.append(
            {
                "key": sector,
                "label": sector,
                "kind": "sector",
                "count": len(rows),
            }
        )
        # Industry subgroups when sector has multiple industries with ≥2 names
        inds = [
            (ind, lst)
            for (sec, ind), lst in by_ind.items()
            if sec == sector and ind and ind != "—" and len(lst) >= 2
        ]
        inds.sort(key=lambda x: (-len(x[1]), x[0]))
        for ind, lst in inds:
            groups.append(
                {
                    "key": f"{sector}||{ind}",
                    "label": f"{sector} › {ind}",
                    "kind": "industry",
                    "count": len(lst),
                }
            )
    return groups


def _parse_group_key(group_key: str) -> tuple[str, str | None]:
    raw = (group_key or "").strip()
    if raw.upper() == ALL_GROUP_KEY:
        return ALL_GROUP_KEY, None
    if "||" in raw:
        sector, industry = raw.split("||", 1)
        return sector.strip(), industry.strip() or None
    return raw, None


def _members_for_group(
    members: list[dict[str, Any]], group_key: str
) -> list[dict[str, Any]]:
    sector, industry = _parse_group_key(group_key)
    if not sector:
        return []
    if sector == ALL_GROUP_KEY:
        return list(members)
    out = [m for m in members if m["sector"] == sector]
    if industry:
        out = [m for m in out if m["industry"] == industry]
    return out


def _expected_as_of_date() -> str:
    """Last completed (or in-progress after open) US equity session date in ET."""
    now = datetime.now(ET)
    d = now.date()
    # Before regular open, prefer previous session for "as of".
    if now.time() < time(9, 30):
        d = d - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d.isoformat()


def _max_bar_date(dated: dict[str, list[tuple[str, float]]], *, skip_spy: bool = False) -> str:
    mx = ""
    for t, series in dated.items():
        if skip_spy and t == "SPY":
            continue
        if series:
            mx = max(mx, series[-1][0])
    return mx


def _ensure_fresh_daily_bars(
    tickers: list[str],
    dated: dict[str, list[tuple[str, float]]],
) -> dict[str, list[tuple[str, float]]]:
    """
    If SPY or group bars lag the expected session date, refresh via market_data
    (IBKR first, Yahoo fallback) and reload from daily_bars.
    """
    expected = _expected_as_of_date()
    member_max = _max_bar_date(dated, skip_spy=True)
    spy_max = (dated.get("SPY") or [("", 0.0)])[-1][0] if dated.get("SPY") else ""
    need_refresh = (not member_max or member_max < expected) or (spy_max < expected)
    if not need_refresh:
        return _append_live_last_day(tickers, dated, expected=expected)

    refresh_list = ["SPY"]
    for t in tickers or []:
        series = dated.get(t) or []
        tmax = series[-1][0] if series else ""
        if not tmax or tmax < expected:
            refresh_list.append(t)
    refresh_list = sorted(set(refresh_list))
    # Always refresh SPY when it lags members or expected.
    if spy_max < expected or (member_max and spy_max < member_max):
        if "SPY" not in refresh_list:
            refresh_list.append("SPY")
            refresh_list = sorted(set(refresh_list))

    try:
        from market_data import load_daily_closes, load_yahoo_daily_closes, preferred_data_source
        from strong_stocks import upsert_daily_bars
    except Exception:
        log.exception("group_movement daily refresh imports failed")
        return dated

    log.info(
        "group_movement refreshing daily bars (member_max=%s spy_max=%s expected=%s n=%s)",
        member_max,
        spy_max,
        expected,
        len(refresh_list),
    )

    # If IBKR is preferred but down, fall back to Yahoo for the whole batch
    # instead of reconnecting (and failing) on every ticker.
    force_yahoo = False
    if preferred_data_source() == "ibkr" and refresh_list:
        probe = refresh_list[0]
        try:
            closes, _hist, meta = load_daily_closes(probe, period="6mo", retries=2)
            if closes is not None and not closes.empty:
                upsert_daily_bars(probe, closes)
            if str((meta or {}).get("provider") or "").startswith("yahoo"):
                force_yahoo = True
            elif (meta or {}).get("ibkr_error"):
                force_yahoo = True
            refresh_list = [t for t in refresh_list if t != probe]
        except Exception:
            force_yahoo = True
            log.exception("probe daily refresh failed for %s", probe)

    for t in refresh_list:
        try:
            if force_yahoo:
                closes, _hist, _meta = load_yahoo_daily_closes(t, period="6mo", retries=2)
            else:
                closes, _hist, _meta = load_daily_closes(t, period="6mo", retries=2)
            if closes is not None and not closes.empty:
                upsert_daily_bars(t, closes)
        except Exception:
            log.exception("refresh daily bars failed for %s", t)
    dated2 = _load_dated_closes(tickers, lookback_calendar_days=max(120, 63 * 4))
    return _append_live_last_day(tickers, dated2, expected=expected)


def _append_live_last_day(
    tickers: list[str],
    dated: dict[str, list[tuple[str, float]]],
    *,
    expected: str | None = None,
) -> dict[str, list[tuple[str, float]]]:
    """
    When the daily history row for today is missing/NaN (common on Yahoo),
    append fast_info lastPrice so the board reaches the current session date.
    """
    expected = expected or _expected_as_of_date()
    if _max_bar_date(dated, skip_spy=True) >= expected:
        # Still fix SPY if only SPY lags.
        spy = dated.get("SPY") or []
        if spy and spy[-1][0] >= expected:
            return dated

    try:
        import pandas as pd
        import yfinance as yf

        from strong_stocks import upsert_daily_bars
    except Exception:
        return dated

    targets = sorted({*(tickers or []), "SPY"})
    wrote = False
    for t in targets:
        series = dated.get(t) or []
        if series and series[-1][0] >= expected:
            continue
        try:
            fi = dict(yf.Ticker(t).fast_info)
            px = fi.get("lastPrice")
            if px is None:
                px = fi.get("last_price")
            px_f = float(px) if px is not None else None
            if px_f is None or px_f <= 0:
                continue
            upsert_daily_bars(t, pd.Series({pd.Timestamp(expected): px_f}))
            wrote = True
        except Exception:
            log.debug("live lastPrice append failed for %s", t, exc_info=True)
    if not wrote:
        return dated
    return _load_dated_closes(tickers, lookback_calendar_days=max(120, 63 * 4))


def _load_dated_closes(
    tickers: list[str], *, lookback_calendar_days: int = 180
) -> dict[str, list[tuple[str, float]]]:
    """ticker → ascending [(date, close), ...]."""
    init_db()
    clean = sorted({(t or "").strip().upper() for t in tickers if t})
    if not clean:
        return {}
    # Include SPY for calendar when available.
    query_tickers = list(clean)
    if "SPY" not in query_tickers:
        query_tickers.append("SPY")
    ph = ",".join("?" * len(query_tickers))
    with get_conn() as conn:
        rows = conn.execute(
            f"""
            SELECT ticker, date, close FROM daily_bars
            WHERE ticker IN ({ph})
              AND date >= date('now', ?)
            ORDER BY ticker ASC, date ASC
            """,
            [*query_tickers, f"-{int(lookback_calendar_days)} days"],
        ).fetchall()
    out: dict[str, list[tuple[str, float]]] = {}
    for r in rows:
        t = (r["ticker"] or "").upper()
        d = str(r["date"] or "")[:10]
        try:
            px = float(r["close"])
        except (TypeError, ValueError):
            continue
        if not d or px <= 0:
            continue
        out.setdefault(t, []).append((d, px))
    return out


def _trading_calendar(
    dated: dict[str, list[tuple[str, float]]], *, window: int
) -> list[str]:
    """
    Last `window` trading dates.

    Prefer SPY when it is as fresh as the group; otherwise build a majority
    calendar from member bars so a stale SPY cannot freeze the board on an
    old As-of date.
    """
    spy = dated.get("SPY") or []
    spy_dates = [d for d, _ in spy]
    member_max = _max_bar_date(dated, skip_spy=True)

    if spy_dates and (not member_max or spy_dates[-1] >= member_max):
        if len(spy_dates) >= window:
            return spy_dates[-window:]
        return list(spy_dates)

    counts: Counter[str] = Counter()
    n_members = 0
    for t, series in dated.items():
        if t == "SPY" or not series:
            continue
        n_members += 1
        # Count only the recent tail to keep the calendar dense.
        for d, _ in series[-(window + 10) :]:
            counts[d] += 1
    if not counts:
        # Absolute fallback: union of everything including SPY.
        all_dates = sorted({d for series in dated.values() for d, _ in series})
        return all_dates[-window:] if len(all_dates) > window else all_dates

    threshold = max(1, n_members // 3)
    ordered = sorted(d for d, c in counts.items() if c >= threshold)
    if len(ordered) < window:
        # Loosen threshold if sparse.
        ordered = sorted(counts.keys())
    if len(ordered) <= window:
        return ordered
    return ordered[-window:]


def _enrich_sessions_from_yahoo(
    tickers: list[str],
    dates: list[str],
    existing: dict[tuple[str, str], dict[str, float | None]],
) -> dict[tuple[str, str], dict[str, float | None]]:
    """
    Fill missing PRE/REGULAR/AFTER via Yahoo prepost intraday (works without TWS).
    Only refreshes tickers that lack recent session rows.
    """
    if not tickers or not dates:
        return existing
    recent = dates[-8:] if len(dates) >= 8 else list(dates)
    latest = dates[-1]
    need: list[str] = []
    for t in tickers:
        latest_cell = existing.get((t, latest)) or {}
        latest_ok = (
            latest_cell.get("pre") is not None
            or latest_cell.get("regular") is not None
            or latest_cell.get("after") is not None
        )
        covered = 0
        for d in recent:
            cell = existing.get((t, d)) or {}
            if (
                cell.get("pre") is not None
                or cell.get("regular") is not None
                or cell.get("after") is not None
            ):
                covered += 1
        # Refresh if latest day is empty, or coverage is thin vs the window.
        if (not latest_ok) or covered < max(3, len(recent) // 2):
            need.append(t)
    if not need:
        return existing

    try:
        from concurrent.futures import ThreadPoolExecutor, as_completed

        from momentum_sessions import refresh_symbol_sessions
    except Exception:
        log.exception("momentum_sessions import failed")
        return existing

    # 15m × 60d covers a 40D/63D research window better than 5m×7d.
    period = "60d" if len(dates) > 10 else "15d"
    log.info("group_movement Yahoo session enrich for %s tickers (%s)", len(need), period)

    def _one(t: str) -> None:
        try:
            refresh_symbol_sessions(t, period=period, interval="15m")
        except Exception:
            log.exception("Yahoo session refresh failed for %s", t)

    workers = min(6, len(need))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futs = [pool.submit(_one, t) for t in need]
        for f in as_completed(futs):
            _ = f.exception()
    return _load_sessions_bulk(tickers, dates)


def _session_provenance(
    tickers: list[str], dates: list[str]
) -> dict[str, dict[str, str]]:
    """
    Per ticker: which provider supplied P/R/A in the window, and latest download time.
    Prefer IBKR when present; otherwise YAHOO; else blank.
    """
    clean = sorted({(t or "").strip().upper() for t in tickers if t})
    days = [str(d)[:10] for d in dates if d]
    if not clean or not days:
        return {}
    try:
        ph_t = ",".join("?" * len(clean))
        ph_d = ",".join("?" * len(days))
        with get_conn() as conn:
            rows = conn.execute(
                f"""
                SELECT symbol, source,
                       MAX(updated_at) AS downloaded_at,
                       COUNT(*) AS n
                FROM momentum_session_obs
                WHERE symbol IN ({ph_t})
                  AND trading_date IN ({ph_d})
                  AND session IN ('PRE', 'REGULAR', 'AFTER')
                  AND return_pct IS NOT NULL
                GROUP BY symbol, source
                """,
                [*clean, *days],
            ).fetchall()
    except Exception:
        log.exception("session provenance load failed")
        return {}

    by_sym: dict[str, dict[str, Any]] = {}
    for r in rows:
        t = (r["symbol"] or "").upper()
        src = (r["source"] or "").upper()
        if src not in ("IBKR", "YAHOO"):
            continue
        slot = by_sym.setdefault(t, {"sources": set(), "downloaded_at": ""})
        slot["sources"].add(src)
        u = str(r["downloaded_at"] or "")
        if u > str(slot["downloaded_at"] or ""):
            slot["downloaded_at"] = u

    out: dict[str, dict[str, str]] = {}
    for t, slot in by_sym.items():
        srcs = slot["sources"]
        if "IBKR" in srcs and "YAHOO" in srcs:
            label = "IBKR"
            note = "IBKR (+Yahoo)"
        elif "IBKR" in srcs:
            label = "IBKR"
            note = "IBKR"
        elif "YAHOO" in srcs:
            label = "YAHOO"
            note = "Yahoo"
        else:
            label = "—"
            note = "—"
        out[t] = {
            "source": label,
            "source_note": note,
            "downloaded_at": str(slot.get("downloaded_at") or ""),
            "downloaded_at_display": _fmt_download_time(slot.get("downloaded_at")),
        }
    return out


def _fmt_download_time(raw: Any) -> str:
    if not raw:
        return "—"
    s = str(raw).strip()
    try:
        if s.endswith("Z"):
            s = s[:-1] + "+00:00"
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(ET).strftime("%Y-%m-%d %H:%M ET")
    except Exception:
        return s[:19] if len(s) >= 19 else s


def _load_sessions_bulk(
    tickers: list[str], dates: list[str]
) -> dict[tuple[str, str], dict[str, float | None]]:
    """
    (ticker, date) → {pre, regular, after} return_pct (percent, not decimal).
    Missing sessions stay None — never fabricate.
    """
    init_db()
    clean = sorted({(t or "").strip().upper() for t in tickers if t})
    days = [str(d)[:10] for d in dates if d]
    if not clean or not days:
        return {}
    # Table may be empty for most NDX names — still query once.
    try:
        ph_t = ",".join("?" * len(clean))
        ph_d = ",".join("?" * len(days))
        with get_conn() as conn:
            # Ensure table exists without importing heavy momentum refresh.
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS momentum_session_obs (
                    symbol TEXT NOT NULL,
                    trading_date TEXT NOT NULL,
                    session TEXT NOT NULL,
                    source TEXT NOT NULL DEFAULT 'YAHOO',
                    start_price REAL,
                    end_price REAL,
                    return_pct REAL,
                    data_status TEXT,
                    updated_at TEXT,
                    PRIMARY KEY (symbol, trading_date, session, source)
                )
                """
            )
            rows = conn.execute(
                f"""
                SELECT symbol, trading_date, session, return_pct, source
                FROM momentum_session_obs
                WHERE symbol IN ({ph_t})
                  AND trading_date IN ({ph_d})
                  AND session IN ('PRE', 'REGULAR', 'AFTER')
                """,
                [*clean, *days],
            ).fetchall()
    except Exception:
        log.exception("group_movement session bulk load failed")
        return {}

    out: dict[tuple[str, str], dict[str, float | None]] = {}
    # Prefer IBKR over YAHOO when both exist for the same cell.
    rank = {"IBKR": 2, "YAHOO": 1}
    seen: dict[tuple[str, str, str], int] = {}
    for r in rows:
        t = (r["symbol"] or "").upper()
        d = str(r["trading_date"] or "")[:10]
        sess = (r["session"] or "").upper()
        src = (r["source"] or "YAHOO").upper()
        sk = (t, d, sess)
        if rank.get(src, 0) < seen.get(sk, 0):
            continue
        seen[sk] = rank.get(src, 0)
        key = (t, d)
        cell = out.setdefault(key, {"pre": None, "regular": None, "after": None})
        try:
            pct = float(r["return_pct"]) if r["return_pct"] is not None else None
        except (TypeError, ValueError):
            pct = None
        if sess == "PRE":
            cell["pre"] = pct
        elif sess == "REGULAR":
            cell["regular"] = pct
        elif sess == "AFTER":
            cell["after"] = pct
    return out


def _persist_session_obs(
    rows: list[dict[str, Any]], *, source: str = "IBKR"
) -> None:
    """Cache session % into momentum_session_obs for reuse (never invent values)."""
    if not rows:
        return
    init_db()
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    payload = []
    for r in rows:
        t = (r.get("ticker") or "").strip().upper()
        d = str(r.get("trading_date") or "")[:10]
        if not t or not d:
            continue
        for sess, key in (
            ("PRE", "pre"),
            ("REGULAR", "regular"),
            ("AFTER", "after"),
        ):
            pct = r.get(key)
            if pct is None:
                continue
            try:
                pct_f = float(pct)
            except (TypeError, ValueError):
                continue
            payload.append(
                {
                    "symbol": t,
                    "trading_date": d,
                    "session": sess,
                    "source": source,
                    "return_pct": pct_f,
                    "data_status": "COMPLETE",
                    "updated_at": now,
                }
            )
    if not payload:
        return
    with get_conn() as conn:
        conn.executemany(
            """
            INSERT INTO momentum_session_obs (
              symbol, trading_date, session, source,
              start_price, end_price, return_pct, data_status, updated_at
            ) VALUES (
              :symbol, :trading_date, :session, :source,
              NULL, NULL, :return_pct, :data_status, :updated_at
            )
            ON CONFLICT(symbol, trading_date, session, source) DO UPDATE SET
              return_pct = excluded.return_pct,
              data_status = excluded.data_status,
              updated_at = excluded.updated_at
            """,
            payload,
        )


def _enrich_sessions_from_ibkr(
    tickers: list[str],
    dates: list[str],
    existing: dict[tuple[str, str], dict[str, float | None]],
) -> dict[tuple[str, str], dict[str, float | None]]:
    """
    Fill missing PRE/REGULAR/AFTER from IBKR extended-hours bars when
    local Full uses IBKR as market data. Merges into existing map.
    """
    try:
        from market_data import preferred_data_source

        if preferred_data_source() != "ibkr":
            return existing
    except Exception:
        return existing
    if not tickers or not dates:
        return existing

    # Prefer IBKR overlay even when Yahoo already filled cells.
    # Skip tickers that already have fresh IBKR rows in the recent window.
    need: list[str] = []
    date_set = set(dates)
    recent = dates[-5:] if len(dates) >= 5 else list(dates)
    ibkr_have: set[str] = set()
    try:
        ph_t = ",".join("?" * len(tickers))
        ph_d = ",".join("?" * len(recent))
        with get_conn() as conn:
            rows = conn.execute(
                f"""
                SELECT DISTINCT symbol
                FROM momentum_session_obs
                WHERE symbol IN ({ph_t})
                  AND trading_date IN ({ph_d})
                  AND session IN ('PRE', 'REGULAR', 'AFTER')
                  AND source = 'IBKR'
                  AND return_pct IS NOT NULL
                """,
                [*tickers, *recent],
            ).fetchall()
        ibkr_have = {(r["symbol"] or "").upper() for r in rows}
    except Exception:
        ibkr_have = set()
    for t in tickers:
        if t not in ibkr_have:
            need.append(t)
    if not need:
        return existing

    try:
        from ibkr_local.adapter import get_adapter

        cal = get_adapter().fetch_session_calendar(
            need, lookback_days=max(12, min(45, len(dates) + 2))
        )
    except Exception:
        log.exception("IBKR session calendar enrich failed")
        return existing

    out = {k: dict(v) for k, v in existing.items()}
    to_store: list[dict[str, Any]] = []
    for t, by_day in (cal or {}).items():
        if not isinstance(by_day, dict):
            continue
        for d, sess in by_day.items():
            if d not in date_set or not isinstance(sess, dict):
                continue
            key = (t, d)
            cell = out.setdefault(key, {"pre": None, "regular": None, "after": None})
            changed = False
            for sk, dk in (("pre", "pre"), ("regular", "regular"), ("after", "after")):
                if sess.get(dk) is None:
                    continue
                try:
                    val = float(sess[dk])
                except (TypeError, ValueError):
                    continue
                # IBKR wins over Yahoo / blanks.
                if cell.get(sk) != val:
                    cell[sk] = val
                    changed = True
            if changed:
                to_store.append(
                    {
                        "ticker": t,
                        "trading_date": d,
                        "pre": cell.get("pre"),
                        "regular": cell.get("regular"),
                        "after": cell.get("after"),
                    }
                )
    try:
        _persist_session_obs(to_store, source="IBKR")
    except Exception:
        log.exception("persist IBKR session obs failed")
    return out


def _ibkr_after_checked(tickers: list[str], trading_date: str) -> set[str]:
    """Tickers that already have an IBKR AFTER observation row for the date (even if null %)."""
    clean = sorted({(t or "").strip().upper() for t in tickers if t})
    d = str(trading_date or "")[:10]
    if not clean or not d:
        return set()
    try:
        ph = ",".join("?" * len(clean))
        with get_conn() as conn:
            rows = conn.execute(
                f"""
                SELECT DISTINCT symbol
                FROM momentum_session_obs
                WHERE symbol IN ({ph})
                  AND trading_date = ?
                  AND session = 'AFTER'
                  AND source = 'IBKR'
                """,
                [*clean, d],
            ).fetchall()
        return {(r["symbol"] or "").upper() for r in rows}
    except Exception:
        return set()


def _refresh_latest_after_from_ibkr(
    tickers: list[str],
    dates: list[str],
    existing: dict[tuple[str, str], dict[str, float | None]],
) -> dict[tuple[str, str], dict[str, float | None]]:
    """
    After RTH, fill AFTER on the latest window day when IBKR has AH bars.
    Skips tickers already checked (IBKR AFTER row present) so quiet AH days
    do not re-query forever.
    """
    try:
        from market_data import preferred_data_source
        from zoneinfo import ZoneInfo

        if preferred_data_source() != "ibkr" or not tickers or not dates:
            return existing
        now_et = datetime.now(ZoneInfo("America/New_York"))
    except Exception:
        return existing
    # Pre-market / early RTH: AH for "today" is not available yet.
    if now_et.hour < 16:
        return existing

    latest = dates[-1]
    checked = _ibkr_after_checked(tickers, latest)
    need = [t for t in tickers if t not in checked]
    if not need:
        return existing

    try:
        from ibkr_local.adapter import get_adapter

        adapter = get_adapter()
        probe = adapter.fetch_session_moves(need[:1], lookback_days=4)
        prow = probe.get(need[0]) if need else None
        if not prow or not prow.get("ok"):
            log.info(
                "skip IBKR latest-after refresh: %s",
                (prow or {}).get("error") if prow else "probe failed",
            )
            return existing
        moves = dict(probe or {})
        rest = need[1:]
        if rest:
            moves.update(adapter.fetch_session_moves(rest, lookback_days=4) or {})
    except Exception:
        log.exception("IBKR latest after refresh failed")
        return existing

    out = {k: dict(v) for k, v in existing.items()}
    to_store: list[dict[str, Any]] = []
    # Persist AFTER row even when pct is None (marker = checked).
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    marker_rows: list[dict[str, Any]] = []
    for t, payload in (moves or {}).items():
        if not isinstance(payload, dict) or not payload.get("ok"):
            continue
        d = str(payload.get("trading_date") or "")[:10] or latest
        if d != latest and d not in dates:
            # Still accept broker's latest session day if it falls in the window.
            if d not in dates:
                continue
        key = (t, d)
        cell = out.setdefault(key, {"pre": None, "regular": None, "after": None})
        for sk, dk in (
            ("pre", "pre_pct"),
            ("regular", "regular_pct"),
            ("after", "after_pct"),
        ):
            if cell.get(sk) is None and payload.get(dk) is not None:
                try:
                    cell[sk] = float(payload[dk])
                except (TypeError, ValueError):
                    pass
        to_store.append(
            {
                "ticker": t,
                "trading_date": d,
                "pre": cell.get("pre"),
                "regular": cell.get("regular"),
                "after": cell.get("after"),
            }
        )
        # Marker so we do not re-poll when AH simply did not trade.
        if cell.get("after") is None and d == latest:
            marker_rows.append(
                {
                    "symbol": t,
                    "trading_date": d,
                    "session": "AFTER",
                    "source": "IBKR",
                    "return_pct": None,
                    "data_status": "NO_BARS",
                    "updated_at": now,
                }
            )
    try:
        _persist_session_obs(to_store, source="IBKR")
        if marker_rows:
            init_db()
            with get_conn() as conn:
                conn.executemany(
                    """
                    INSERT INTO momentum_session_obs (
                      symbol, trading_date, session, source,
                      start_price, end_price, return_pct, data_status, updated_at
                    ) VALUES (
                      :symbol, :trading_date, :session, :source,
                      NULL, NULL, :return_pct, :data_status, :updated_at
                    )
                    ON CONFLICT(symbol, trading_date, session, source) DO UPDATE SET
                      data_status = excluded.data_status,
                      updated_at = excluded.updated_at
                    """,
                    marker_rows,
                )
    except Exception:
        log.exception("persist latest after refresh failed")
    return out


def _close_map(series: list[tuple[str, float]]) -> dict[str, float]:
    return {d: px for d, px in series}


def _build_row(
    member: dict[str, Any],
    *,
    dates: list[str],
    dated: dict[str, list[tuple[str, float]]],
    sessions: dict[tuple[str, str], dict[str, float | None]],
) -> dict[str, Any]:
    t = member["ticker"]
    cmap = _close_map(dated.get(t) or [])
    # Need prior close before window for day-0 daily % when possible.
    series = dated.get(t) or []
    prior_by_date: dict[str, float] = {}
    for i, (d, px) in enumerate(series):
        if i > 0:
            prior_by_date[d] = series[i - 1][1]

    days_out: list[dict[str, Any]] = []
    for i, d in enumerate(dates):
        close = cmap.get(d)
        prior = prior_by_date.get(d)
        daily_pct: float | None = None
        if close is not None and prior is not None and prior > 0:
            daily_pct = (close / prior - 1.0) * 100.0

        sess = sessions.get((t, d)) or {}

        def _round_pct(v: Any) -> float | None:
            if v is None:
                return None
            try:
                return round(float(v), 2)
            except (TypeError, ValueError):
                return None

        days_out.append(
            {
                "date": d,
                "daily_pct": round(daily_pct, 2) if daily_pct is not None else None,
                "pre_pct": _round_pct(sess.get("pre")),
                "regular_pct": _round_pct(sess.get("regular")),
                "after_pct": _round_pct(sess.get("after")),
                "close": round(close, 4) if close is not None else None,
            }
        )

    last_day = days_out[-1] if days_out else {}
    return {
        "ticker": t,
        "name": member.get("name") or "",
        "sector": member.get("sector") or UNCLASSIFIED,
        "industry": member.get("industry") or "—",
        "days": days_out,
        "last_daily_pct": last_day.get("daily_pct"),
        "data_source": "—",
        "data_source_note": "—",
        "downloaded_at": "",
        "downloaded_at_display": "—",
    }


def load_group_movement(
    *,
    group_key: str | None = None,
    window: int = DEFAULT_WINDOW,
    refresh: bool = False,
) -> dict[str, Any]:
    """
    Build the Group Movement research payload for one selected group.

    refresh=False (default): read daily_bars + cached session rows only — fast.
    refresh=True: pull fresh daily closes and P/R/A for the *current* group only.
    """
    try:
        w = int(window)
    except (TypeError, ValueError):
        w = DEFAULT_WINDOW
    if w not in WINDOW_CHOICES:
        # Nearest allowed
        w = min(WINDOW_CHOICES, key=lambda x: abs(x - w))

    members = list_ndx100_members()
    groups = list_movement_groups(members)
    if not group_key:
        group_key = groups[0]["key"] if groups else UNCLASSIFIED
    selected = _members_for_group(members, group_key)
    # Fallback if key stale
    if not selected and groups:
        group_key = groups[0]["key"]
        selected = _members_for_group(members, group_key)

    tickers = [m["ticker"] for m in selected]
    dated = _load_dated_closes(tickers, lookback_calendar_days=max(120, w * 4))
    if refresh:
        dated = _ensure_fresh_daily_bars(tickers, dated)
    dates = _trading_calendar(dated, window=w)
    sessions = _load_sessions_bulk(tickers, dates) if dates else {}
    is_all = (group_key or "").strip().upper() == ALL_GROUP_KEY
    if refresh and dates and tickers:
        # Explicit refresh for this group only (not the whole Research universe).
        sessions = _enrich_sessions_from_yahoo(tickers, dates, sessions)
        # IBKR overlay: sector/industry sized pools. ALL stays Yahoo/cache to avoid
        # multi-minute freezes (~100 names × historical bars).
        if not is_all:
            sessions = _enrich_sessions_from_ibkr(tickers, dates, sessions)
            sessions = _refresh_latest_after_from_ibkr(tickers, dates, sessions)

    rows = [
        _build_row(m, dates=dates, dated=dated, sessions=sessions) for m in selected
    ]
    prov = _session_provenance(tickers, dates) if dates else {}
    for r in rows:
        meta = prov.get(r["ticker"]) or {}
        r["data_source"] = meta.get("source") or "—"
        r["data_source_note"] = meta.get("source_note") or r["data_source"]
        r["downloaded_at"] = meta.get("downloaded_at") or ""
        r["downloaded_at_display"] = meta.get("downloaded_at_display") or "—"
    # ALL: cluster by sector/industry so Health Care etc. are easy to find.
    # Single-sector views: keep latest daily % desc.
    if (group_key or "").strip().upper() == ALL_GROUP_KEY:
        rows.sort(
            key=lambda r: (
                (r.get("sector") or "").lower(),
                (r.get("industry") or "").lower(),
                -(
                    r["last_daily_pct"]
                    if r.get("last_daily_pct") is not None
                    else float("-inf")
                ),
                r.get("ticker") or "",
            )
        )
    else:
        rows.sort(
            key=lambda r: (
                -(
                    r["last_daily_pct"]
                    if r.get("last_daily_pct") is not None
                    else float("-inf")
                ),
                r.get("ticker") or "",
            )
        )

    sector, industry = _parse_group_key(group_key or "")
    date_labels = []
    for d in dates:
        # Compact header: "Aug 1" style via month day if parseable
        try:
            y, m, day = d.split("-")
            months = (
                "Jan",
                "Feb",
                "Mar",
                "Apr",
                "May",
                "Jun",
                "Jul",
                "Aug",
                "Sep",
                "Oct",
                "Nov",
                "Dec",
            )
            label = f"{months[int(m) - 1]} {int(day)}"
        except Exception:
            label = d
        date_labels.append({"date": d, "label": label})

    return {
        "pool": "NASDAQ-100",
        "window": w,
        "window_choices": list(WINDOW_CHOICES),
        "group_key": group_key,
        "group_sector": sector,
        "group_industry": industry,
        "groups": groups,
        "dates": dates,
        "date_labels": date_labels,
        "rows": rows,
        "row_count": len(rows),
        "member_count": len(members),
        "columns": get_column_prefs(),
        "as_of": dates[-1] if dates else "",
        "refreshed": bool(refresh),
        "notes": (
            "Daily Change % by date. Premarket/Regular/After-hours from IBKR when "
            "connected, otherwise Yahoo. No Leader/Follower labels. "
            "Use Refresh Group to update the current group only."
        ),
    }


def export_group_movement_csv(
    *,
    group_key: str | None = None,
    window: int = DEFAULT_WINDOW,
) -> tuple[str, str]:
    """
    Build CSV text for the current Group Movement view.
    Returns (filename, csv_text).
    """
    import csv
    import io

    payload = load_group_movement(group_key=group_key, window=window)
    dates = list(payload.get("dates") or [])
    rows = list(payload.get("rows") or [])
    gk = (payload.get("group_key") or "group").replace("|", "-").replace(" ", "_")
    as_of = payload.get("as_of") or "na"
    fname = f"group_movement_{gk}_{payload.get('window') or window}D_{as_of}.csv"

    buf = io.StringIO()
    # Excel (many locales) opens UTF-16 LE TSV more reliably than UTF-8 CSV.
    w = csv.writer(buf, delimiter="\t", lineterminator="\r\n")
    header = [
        "Ticker",
        "Name",
        "Sector",
        "Industry",
        "Source",
        "Downloaded",
    ]
    for d in dates:
        header.extend(
            [
                f"{d}_daily_pct",
                f"{d}_pre_pct",
                f"{d}_regular_pct",
                f"{d}_after_pct",
                f"{d}_close",
            ]
        )
    w.writerow(header)

    def _fmt(v: Any) -> str:
        if v is None:
            return ""
        return str(v)

    for r in rows:
        line = [
            r.get("ticker") or "",
            r.get("name") or "",
            r.get("sector") or "",
            r.get("industry") or "",
            r.get("data_source") or "",
            r.get("downloaded_at_display") or r.get("downloaded_at") or "",
        ]
        by_date = {d.get("date"): d for d in (r.get("days") or []) if d.get("date")}
        for d in dates:
            cell = by_date.get(d) or {}
            line.extend(
                [
                    _fmt(cell.get("daily_pct")),
                    _fmt(cell.get("pre_pct")),
                    _fmt(cell.get("regular_pct")),
                    _fmt(cell.get("after_pct")),
                    _fmt(cell.get("close")),
                ]
            )
        w.writerow(line)

    return fname, buf.getvalue()


def export_group_movement_xlsx(
    *,
    group_key: str | None = None,
    window: int = DEFAULT_WINDOW,
) -> tuple[str, bytes]:
    """Build an Excel workbook (.xlsx) for Group Movement. Returns (filename, bytes)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    payload = load_group_movement(group_key=group_key, window=window)
    dates = list(payload.get("dates") or [])
    rows = list(payload.get("rows") or [])
    gk = (payload.get("group_key") or "group").replace("|", "-").replace(" ", "_")
    as_of = payload.get("as_of") or "na"
    fname = f"group_movement_{gk}_{payload.get('window') or window}D_{as_of}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Group Movement"

    header = [
        "Ticker",
        "Name",
        "Sector",
        "Industry",
        "Source",
        "Downloaded",
    ]
    for d in dates:
        header.extend(
            [
                f"{d} Daily %",
                f"{d} Pre %",
                f"{d} Regular %",
                f"{d} After %",
                f"{d} Close",
            ]
        )
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        line: list[Any] = [
            r.get("ticker") or "",
            r.get("name") or "",
            r.get("sector") or "",
            r.get("industry") or "",
            r.get("data_source") or "",
            r.get("downloaded_at_display") or r.get("downloaded_at") or "",
        ]
        by_date = {d.get("date"): d for d in (r.get("days") or []) if d.get("date")}
        for d in dates:
            cell = by_date.get(d) or {}
            line.extend(
                [
                    cell.get("daily_pct"),
                    cell.get("pre_pct"),
                    cell.get("regular_pct"),
                    cell.get("after_pct"),
                    cell.get("close"),
                ]
            )
        ws.append(line)

    # Reasonable widths for identity columns
    widths = (10, 28, 22, 28, 10, 22)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"  # freeze header only — more compatible with Excel/WPS

    bio = io.BytesIO()
    wb.save(bio)
    return fname, bio.getvalue()
